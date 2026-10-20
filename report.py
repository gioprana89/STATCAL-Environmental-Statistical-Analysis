"""Ekspor Excel bertabel dan berformat, menggunakan openpyxl pada aplikasi pengguna."""
from __future__ import annotations
from datetime import datetime
from io import BytesIO
import re
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from stat_core import Result

HEAD='164453';ACCENT='0F827A';LIGHT='E7F3F1';GRID='E1EAED'

def _cell(x):
    if x is None or x is pd.NA:return None
    if isinstance(x,np.generic):x=x.item()
    if isinstance(x,(float,np.floating)) and (np.isnan(x) or np.isinf(x)):return None
    if isinstance(x,(pd.Timestamp,datetime)):return x.to_pydatetime() if isinstance(x,pd.Timestamp) else x
    if isinstance(x, (pd.Timedelta,)):return str(x)
    if isinstance(x,str) and x.startswith(('=','+','-','@')):return "'"+x  # cegah formula injection
    return x

def _sheet(wb,name):
    name=re.sub(r'[:\\/*?\[\]]','_',name)[:31]
    base=name;counter=1
    while name in wb.sheetnames:
        postfix=f'_{counter}';name=base[:31-len(postfix)]+postfix;counter+=1
    return wb.create_sheet(name)

def _format(ws,title,columns,width_limit=42):
    ws.sheet_view.showGridLines=False
    ws.freeze_panes='A5';ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.page_margins=PageMargins(left=.3,right=.3,top=.5,bottom=.5,header=.2,footer=.2)
    n=max(len(columns),2)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
    h=ws.cell(1,1,title);h.font=Font(name='Aptos Display',size=15,bold=True,color='FFFFFF')
    h.fill=PatternFill('solid',fgColor=HEAD);h.alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=32
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n)
    p=ws.cell(2,1,'STATCAL ONLINE — Environmental Statistical Analysis | Data dan hasil bersifat sesuai sumber yang dipilih pengguna')
    p.font=Font(size=9,color=ACCENT,italic=True);ws.row_dimensions[2].height=19
    for i,col in enumerate(columns,1):
        c=ws.cell(4,i,str(col));c.fill=PatternFill('solid',fgColor=ACCENT)
        c.font=Font(name='Aptos',bold=True,color='FFFFFF',size=10)
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width=min(width_limit,max(14,min(34,len(str(col))+5)))
    ws.row_dimensions[4].height=35
    ws.auto_filter.ref=f'A4:{get_column_letter(len(columns))}{ws.max_row}'

def _add_df(wb,name,title,frame):
    ws=_sheet(wb,name)
    df=frame.copy()
    if not isinstance(df,pd.DataFrame): df=pd.DataFrame(df)
    df.columns=[str(c) for c in df.columns]
    _format(ws,title,list(df.columns) or ['Keterangan'])
    if df.empty:
        ws.cell(5,1,'Tidak ada hasil yang dapat ditampilkan.')
    else:
        for r,values in enumerate(df.itertuples(index=False,name=None),5):
            for j,v in enumerate(values,1):
                cell=ws.cell(r,j,_cell(v));cell.font=Font(name='Aptos',size=10,color=HEAD)
                cell.alignment=Alignment(vertical='center')
                if r%2==0:cell.fill=PatternFill('solid',fgColor='F1F8F8')
                if isinstance(v,(float,np.floating)):
                    cell.number_format='0.000000'
                if isinstance(v,(int,np.integer)):
                    cell.number_format='#,##0'
                cell.border=Border(bottom=Side(style='hair',color=GRID))
    ws.auto_filter.ref=f'A4:{get_column_letter(max(len(df.columns),1))}{max(5,4+len(df))}'
    ws.print_title_rows='1:4'
    return ws


def report_bytes(results: list[tuple[str,Result]], selected: pd.DataFrame,
                 title: str, source: str, options: dict) -> bytes:
    wb=Workbook(); wb.remove(wb.active)
    metadata=[('Judul analisis',title),('Sumber data',source),('Jumlah baris setelah filter',len(selected)),
              ('Waktu penyusunan laporan','Dihasilkan saat pengguna mengunduh laporan'),
              ('Peringatan','Data simulasi hanya untuk pembelajaran, tidak untuk kesimpulan empiris.')]
    metadata.extend((str(k),str(v)) for k,v in options.items())
    _add_df(wb,'INFORMASI','INFORMASI ANALISIS',pd.DataFrame(metadata,columns=['Parameter','Nilai']))
    # gabungkan tabel yang sejenis agar jumlah sheet tetap wajar
    names=[]
    for group,res in results:
        for name in res.tables:
            if name not in names:names.append(name)
    for name in names:
        all_rows=[]
        for group,res in results:
            frame=res.tables.get(name)
            if frame is None:continue
            block=frame.copy();block.insert(0,'Grup / Split',group)
            all_rows.append(block)
        if all_rows:
            whole=pd.concat(all_rows,ignore_index=True)
            _add_df(wb,name[:27],f'{title} — {name.replace("_"," ")}',whole)
    note_rows=[]
    for group,res in results:
        for note in res.note:note_rows.append({'Grup / Split':group,'Keterangan':note})
    if note_rows:_add_df(wb,'CATATAN','CATATAN METODOLOGI',pd.DataFrame(note_rows))
    if len(selected)<=50000:
        _add_df(wb,'DATA_FILTER','DATA SETELAH FILTER',selected)
    else:
        _add_df(wb,'DATA_FILTER','CATATAN DATA',pd.DataFrame([{'Informasi':'Dataset > 50.000 baris; tidak diekspor untuk menghemat sumber daya server.'}]))
    bio=BytesIO();wb.save(bio);return bio.getvalue()
