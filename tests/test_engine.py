from __future__ import annotations
import sys
from pathlib import Path
import numpy as np
import pandas as pd
import pytest
from scipy import stats
from openpyxl import load_workbook
from io import BytesIO
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from data_io import parse,category_cols,apply_filters,detect_header
from stat_core import corr_pair,corr_matrix,regression,paired_values,test_difference as difference_test,grouped,_dunn_posthoc
from plots import correlation_png,scatter_png,regression_png,qq_png,difference_png
from report import report_bytes

@pytest.fixture
def demo():return parse((ROOT/'data/Kasus_1_Sepeda_Motor_CO.csv').read_bytes(),'demo.csv')
@pytest.fixture
def air():return parse((ROOT/'data/Kasus_2_Kualitas_Air_Tiga_Lokasi.csv').read_bytes(),'air.csv')

def test_csv_size(demo):assert demo.shape==(48,11)
def test_group_sizes(demo):assert [len(block) for _,block in grouped(demo,['Jenis_hari'])]==[24,24]
def test_two_way_split(demo):assert sorted(len(d) for _,d in grouped(demo,['Tahun','Jenis_hari']))==[12]*4
def test_filter(demo):assert len(apply_filters(demo,{'Tahun':[2024],'Jenis_hari':['Hari kerja']}))==12
def test_numeric(demo):assert demo['CO_rata2_ppm'].dtype.kind=='f'
def test_correlation(demo):
 r,xy=corr_pair(demo,'Motor_rata2_15menit','CO_rata2_ppm')
 assert r['N pasangan']==48 and abs(r['Koefisien']-.8028771576946079)<1e-10
 assert r['CI 95% bawah']<r['Koefisien']<r['CI 95% atas']
@pytest.mark.parametrize('method',['pearson','spearman','kendall'])
def test_corr_methods(demo,method):
 r,_=corr_pair(demo,'Motor_rata2_15menit','CO_rata2_ppm',method)
 assert 0<r['Koefisien']<=1 and r['p-value']<.05

def test_corr_matrix(demo):
 r=corr_matrix(demo,['Motor_rata2_15menit','CO_rata2_ppm','Suhu_rata2_C'])
 assert len(r.tables['Pasangan'])==3
 assert all(np.diag(r.meta['matrix'])==1)

def test_corr_fails_constant(demo):
 with pytest.raises(ValueError):corr_pair(demo,'N_hari_sampel','CO_rata2_ppm')

def test_corr_missing():
 d=pd.DataFrame({'X':[1.,2.,3.,4.,np.nan],'Y':[2.,3.,4.,6.,2.]})
 r,_=corr_pair(d,'X','Y');assert r['N pasangan']==4

def test_reg_simple(demo):
 r=regression(demo,'CO_rata2_ppm',['Motor_rata2_15menit'])
 assert abs(r.tables['Ringkasan_Model']['R Square'].iloc[0]-.8028771576946079**2)<1e-9
 assert len(r.tables['Koefisien'])==2

def test_reg_multiple(demo):
 r=regression(demo,'CO_rata2_ppm',['Motor_rata2_15menit','Suhu_rata2_C','Angin_rata2_mps'])
 assert len(r.tables['Koefisien'])==4
 assert 'VIF' in r.tables and 'Diagnostik' in r.tables
 assert r.tables['Ringkasan_Model']['N valid'].iloc[0]==48

def test_reg_robust(demo):
 plain=regression(demo,'CO_rata2_ppm',['Motor_rata2_15menit'])
 robust=regression(demo,'CO_rata2_ppm',['Motor_rata2_15menit'],robust=True)
 assert np.allclose(plain.tables['Koefisien']['Koefisien'],robust.tables['Koefisien']['Koefisien'])
 assert robust.tables['Ringkasan_Model']['SE'].iloc[0]=='HC3 robust'

def test_reg_constant_rejected(demo):
 with pytest.raises(ValueError):regression(demo,'CO_rata2_ppm',['N_hari_sampel'])

def test_paired_count(demo):
 paired=paired_values(demo,'Jenis_hari','CO_rata2_ppm','Hari kerja','Akhir pekan','Periode')
 assert len(paired)==24

def test_paired_duplicate_rejected(demo):
 bad=pd.concat([demo,demo.iloc[[0]]],ignore_index=True)
 with pytest.raises(ValueError,match='lebih dari satu'):paired_values(bad,'Jenis_hari','CO_rata2_ppm','Hari kerja','Akhir pekan','Periode')

@pytest.mark.parametrize('method',['paired_t','wilcoxon','independent_t','welch_t','mannwhitney'])
def test_tests_two_groups(demo,method):
 r=difference_test(demo,method,'CO_rata2_ppm','Jenis_hari','Hari kerja','Akhir pekan','Periode')
 assert r.tables['Hasil_Uji']['p-value'].iloc[0]<.001

@pytest.mark.parametrize('method',['anova','welch_anova','kruskal'])
def test_tests_three_groups(air,method):
 r=difference_test(air,method,'BOD_mg_L','Lokasi',posthoc=True)
 assert r.tables['Hasil_Uji']['p-value'].iloc[0]<.01
 assert any(x.startswith('Posthoc') for x in r.tables)


def test_independent_data_by_year(demo):
 for key,sub in grouped(demo,['Tahun']):
  assert len(sub)==24
  r=difference_test(sub,'paired_t','CO_rata2_ppm','Jenis_hari','Hari kerja','Akhir pekan','Periode')
  assert r.tables['Hasil_Uji']['N pasangan'].iloc[0]==12

def test_wide_paired(demo):
 wide=demo.pivot(index='Periode',columns='Jenis_hari',values='CO_rata2_ppm').reset_index()
 r=difference_test(wide,'wilcoxon','Hari kerja',response_b='Akhir pekan')
 assert r.tables['Hasil_Uji']['N pasangan'].iloc[0]==24

def test_png_all(demo,air):
 cm=corr_matrix(demo,['Motor_rata2_15menit','CO_rata2_ppm'])
 reg=regression(demo,'CO_rata2_ppm',['Motor_rata2_15menit'])
 ts=difference_test(demo,'paired_t','CO_rata2_ppm','Jenis_hari','Hari kerja','Akhir pekan','Periode')
 dat=demo[['Motor_rata2_15menit','CO_rata2_ppm']]
 arr=[correlation_png(cm,'Korelasi','Simulasi'),
      scatter_png(dat,'Motor_rata2_15menit','CO_rata2_ppm','Scatter','Simulasi'),
      regression_png(reg,'Regresi','Simulasi'),qq_png(reg,'QQ','Simulasi'),
      difference_png(ts,'Uji Beda','Simulasi')]
 for b in arr:assert b[:8]==b'\x89PNG\r\n\x1a\n' and len(b)>3000

def test_excel_report(demo):
 reg=regression(demo,'CO_rata2_ppm',['Motor_rata2_15menit'])
 r=report_bytes([('Semua data',reg)],demo,'Contoh Hasil Statistik','Simulasi',{'alpha':.05})
 wb=load_workbook(BytesIO(r),read_only=True)
 assert 'Ringkasan_Model' in wb.sheetnames and 'Koefisien' in wb.sheetnames
 assert 'INFORMASI' in wb.sheetnames and 'DATA_FILTER' in wb.sheetnames
 assert wb['Koefisien']['A1'].value=='Contoh Hasil Statistik — Koefisien'

def test_auto_header_xlsx():
 import pandas as pd
 raw=(ROOT/'data/Kasus_1_Sepeda_Motor_CO.xlsx').read_bytes()
 sheet=pd.ExcelFile(BytesIO(raw)).sheet_names[0]
 assert detect_header(raw,sheet)>0


def test_posthoc_ties():
 g={'A':np.array([1,2,2,3]),'B':np.array([3,4,5,5]),'C':np.array([6,7,7,8])}
 out=_dunn_posthoc(g)
 assert len(out)==3 and out['p-value (Holm)'].between(0,1).all()
